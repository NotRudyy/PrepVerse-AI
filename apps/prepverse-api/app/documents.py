from io import BytesIO


def extract_text(filename: str, content: bytes) -> str:
    suffix = filename.rsplit(".", 1)[-1].lower()
    if suffix == "pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)
    if suffix == "pptx":
        from pptx import Presentation

        return "\n".join(shape.text for slide in Presentation(BytesIO(content)).slides for shape in slide.shapes if hasattr(shape, "text"))
    if suffix == "docx":
        from docx import Document

        document = Document(BytesIO(content))
        pieces = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            pieces.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(pieces)
    raise ValueError("Only PDF, PPTX, and DOCX study materials are supported.")
